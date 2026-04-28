import fitz
import docx
import os


def extract_text(file_path):
    extension=os.path.splitext(file_path)[1].lower()

    if extension=='.pdf':
        return _extract_pdf(file_path)

    elif extension in ['.docx','.doc']:
        return _extract_docx(file_path)

    elif extension=='.txt':
        return _extract_txt(file_path)

    else:
        raise ValueError("Unsupported file format")


def _extract_pdf(file_path):
    text=""
    try:
        doc=fitz.open(file_path)

        for page in doc:
            text+=page.get_text("text")+"\n"

        doc.close()

    except Exception as e:
        print("PDF extraction error:",e)

    return text


def _extract_docx(file_path):
    text=""

    try:
        doc=docx.Document(file_path)

        for para in doc.paragraphs:
            text+=para.text+"\n"
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text+=cell.text+" "

    except Exception as e:
        print("DOCX extraction error:",e)

    return text


def _extract_txt(file_path):
    try:
        with open(file_path,'r',encoding='utf-8',errors='ignore') as f:
            return f.read()

    except Exception as e:
        print("TXT extraction error:",e)
        return ""